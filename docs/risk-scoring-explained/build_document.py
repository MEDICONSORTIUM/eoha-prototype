from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

out = Path('docs/risk-scoring-explained')
doc = Document()
s = doc.sections[0]
s.page_height = Inches(11.7); s.page_width = Inches(8.3)
s.top_margin = s.bottom_margin = Inches(.7)
s.left_margin = s.right_margin = Inches(.8)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2']:
    st = doc.styles[name]; st.font.name = 'Calibri'; st.font.color.rgb = RGBColor(0,0,0)
doc.styles['Normal'].font.size = Pt(11)
doc.styles['Normal'].paragraph_format.space_after = Pt(8)
doc.styles['Normal'].paragraph_format.line_spacing = 1.08
doc.styles['Title'].font.size = Pt(25)
doc.styles['Heading 1'].font.size = Pt(17)
doc.styles['Heading 2'].font.size = Pt(12)
md=[]
def p(t):
    doc.add_paragraph(t); md.append(t+'\n')
def h(t):
    doc.add_heading(t,1); md.append('## '+t+'\n')
def sub(t):
    doc.add_heading(t,2); md.append('### '+t+'\n')
def page(): doc.add_page_break()
def table(headers, rows, widths):
    t=doc.add_table(rows=1, cols=len(headers)); t.autofit=False
    for c,w in zip(t.columns,widths): c.width=Inches(w)
    for c,x in zip(t.rows[0].cells,headers): c.text=x
    for row in rows:
        for c,x in zip(t.add_row().cells,row): c.text=x
    for i,row in enumerate(t.rows):
        for c,w in zip(row.cells,widths):
            c.width=Inches(w)
            pr=c._tc.get_or_add_tcPr()
            borders=OxmlElement('w:tcBorders')
            for side in ['top','left','bottom','right']:
                e=OxmlElement('w:'+side); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'4'); e.set(qn('w:color'),'D9D9D9'); borders.append(e)
            pr.append(borders)
            margins=OxmlElement('w:tcMar')
            for side in ['top','left','bottom','right']:
                e=OxmlElement('w:'+side); e.set(qn('w:w'),'90'); e.set(qn('w:type'),'dxa'); margins.append(e)
            pr.append(margins)
            if i==0:
                e=OxmlElement('w:shd'); e.set(qn('w:fill'),'E8EEF3'); pr.append(e)
            for para in c.paragraphs:
                para.paragraph_format.space_after=Pt(3)
                for r in para.runs: r.font.size=Pt(10); r.bold=(i==0)
        keep=OxmlElement('w:cantSplit'); row._tr.get_or_add_trPr().append(keep)
    t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    md.append('| '+' | '.join(headers)+' |\n| '+' | '.join(['---']*len(headers))+' |\n'+'\n'.join('| '+' | '.join(r)+' |' for r in rows)+'\n')

doc.add_paragraph('EOHA risk scoring explained','Title'); md.append('# EOHA risk scoring explained\n')
p('Method, rationale and sources | 27 October 2026')
p('EOHA currently calculates a malaria environmental screening score using four fixed rules. The score summarises conditions that may be relevant to malaria surveillance at ward level. It is a prototype index from 0 to 100 points, not a calibrated probability of infection or an outbreak. The NCD module has a separate proposed exposure framework, but its scoring remains disabled.')
h('How the malaria score is calculated')
p('For each ward observation, EOHA adds points for soil moisture, land surface temperature, a water index and population density. The active dashboard implementation is calculate_risk_score in scripts/malaria/dashboard/risk.js; the reports page repeats these rules in calculateRisk in scripts/malaria/reports.js. [P1, P2]')
table(['Input','Rule applied to a valid value','Points'],[
('Soil moisture','Above 0.35 m³/m³','40'),
('Soil moisture','Above 0.25 and up to 0.35 m³/m³','20'),
('Soil moisture','0.25 m³/m³ or below','0'),
('Surface temperature','25°C to 30°C inclusive; otherwise zero','30'),
('NDWI water index','Above −0.1; otherwise zero','20'),
('Population density','Above 300 people/km²; otherwise zero','10')],[1.45,4.25,.9])
p('Total score = moisture points + temperature points + water points + population points. Only one soil-moisture band is counted. The maximum is 40 + 30 + 20 + 10 = 100 points.')
p('The exact boundaries matter: moisture of 0.35 receives 20 points; moisture of 0.25 receives zero. NDWI of −0.1 and density of 300 each receive zero. Temperatures of exactly 25°C and 30°C both receive 30 points.')
sub('How the score is labelled')
p('Low: below 25 points. Moderate: at least 25 and below 50 points. High: at least 50 points. These are project classification bands, not clinical thresholds. Because the components are multiples of 10, individual complete ward records produce scores in steps of 10; averages can fall between those steps. [P1, P2]')
page()
h('Why this approach is used')
p('A fixed-rule index provides a simple, explainable baseline for the prototype: each contribution can be inspected, calculations can run in the browser, and the same rules can drive map colours and report filters. It supports exploration of where selected environmental conditions coincide and where further surveillance may be useful. This is a practical rationale for the implementation, not evidence that its chosen weights predict disease accurately.')
sub('Why these four inputs are relevant')
p('Soil moisture is used as a proxy for wet environmental conditions that may support vector habitat. WHO describes temperature, rainfall and humidity as influences on mosquito reproduction and activity. That supports examining environmental conditions, but does not establish EOHA’s 0.25 and 0.35 soil-moisture cutoffs or its 40-point maximum. Soil moisture is not a direct measurement of standing breeding water. [1]')
p('Temperature affects mosquito and parasite development. Mordecai and colleagues reported a modelled transmission optimum around 25°C and declining transmission above about 28°C. This supports temperature as a predictor, but does not validate awarding identical points throughout 25–30°C. EOHA uses satellite surface temperature; the biological temperature relationships cannot be transferred directly to this measurement without validation. [2, 3]')
p('NDWI, the Normalized Difference Water Index, uses green and near-infrared reflectance to help identify open water. The exporter calculates (green − near infrared) / (green + near infrared). Water-related information is relevant to habitat screening, but a ward-average NDWI above −0.1 does not establish breeding-site presence or infected mosquitoes. The exact threshold and 20 points are project rules. [P3, 4]')
p('Population density adds a proxy for the concentration of people potentially exposed. It does not directly measure exposure, infection, immunity or access to prevention. The >300 people/km² cutoff and 10 points should be treated as project assumptions; a higher density is not established here as a universally higher malaria transmission risk. [P1, 5]')
sub('Worked example')
p('Consider a hypothetical ward with soil moisture of 0.30 m³/m³, surface temperature of 27°C, NDWI of 0.05 and population density of 200 people/km². Its score is 20 + 30 + 20 + 0 = 70 points, labelled High under EOHA’s rules. This means the ward meets several selected screening conditions. It does not mean a resident has a 70% chance of malaria.')
p('The dashboard helper calculates a rounded arithmetic mean of ward scores for a group, with equal weight per ward, and returns unavailable for an empty group. This is not a population-weighted infection estimate. Some interface elements show a percent sign; the scientifically appropriate reading of the current calculation is points out of 100. [P1]')
page()
h('Where the data and scoring rules come from')
p('There are three distinct sources: the project code defines the numerical rules; environmental products supply measurements; scientific publications explain why environmental indicators may matter. A dataset provider or relevant paper should not be cited as the origin of the complete EOHA formula unless it actually specifies that formula.')
sub('Origin of the numerical rules')
p('The exact weights of 40, 30, 20 and 10 points, the input cutoffs, and the Low/Moderate/High bands are documented in the repository implementation. Their published derivation is not established in the project documentation. They should therefore be described as project-defined heuristic rules, pending calibration and validation, rather than a WHO-approved model or a formula reproduced from a specific research paper. [P1, P2]')
sub('Environmental products identified in the exporter')
table(['Field','Product and processing in the code'],[
('Surface temperature','NASA MODIS/061/MOD11A1, LST_Day_1km. Daily values are averaged by month, scaled by 0.02 and converted from kelvin to Celsius. [3]'),
('Soil moisture','ECMWF/ERA5_LAND/HOURLY, volumetric_soil_water_layer_1. Hourly values are averaged by month. This is reanalysis, not a direct satellite observation. [6]'),
('NDWI','NASA MODIS/061/MOD09GA, green band sur_refl_b04 and near-infrared band sur_refl_b02. The index is calculated for images and averaged by month. [4]'),
('Population density','WorldPop/GP/100m/pop, 2020 population. The exporter multiplies the ward mean of pixel counts by 100 as an approximate density conversion; this requires spatial-area validation. [5]')],[1.45,5.15])
p('The exporter reduces environmental images over ward boundaries at a requested 1,000 m scale. That does not make every input a native 1 km measurement. In particular, ERA5-Land is substantially coarser. The exporter also includes air temperature, vegetation and land-cover fields, but these do not contribute to the current four-term malaria score. [P3, 6]')
p('These product mappings describe the current exporter. They do not independently certify the provenance of every row in the existing CSV: a dated export manifest and processing history are needed for that. The WorldPop collection referenced by the code is now marked deprecated in the catalogue, and its 2020 population baseline should be explicit when used with later observations. [P3, 5]')
page()
h('Limitations and the NCD scoring proposal')
sub('Limits of the malaria implementation')
p('The index combines correlated moisture and water signals, uses abrupt thresholds, and omits confirmed malaria cases, vector species, intervention coverage, immunity and movement of infected people. Its components can compensate for each other: a high total is possible even when temperature contributes no points. Consequently, it should not be interpreted as a complete transmission model.')
p('Missing-data handling needs correction before operational use. The dashboard converts missing or invalid numbers to zero; for NDWI, zero exceeds −0.1 and therefore earns 20 points. The reports implementation compares raw parsed values, so coercion can also produce misleading results. Missing inputs should be flagged and the score withheld or handled under an explicit reviewed policy. [P1, P2]')
p('The exporter does not explicitly apply the available MODIS quality-bit filters before monthly averaging. Cloud screening, valid-observation coverage, population-density conversion and consistency between surface and air temperature require review. Monthly averages also hide short-lived conditions and do not by themselves establish a forecast lead time. [P3, 3, 4]')
p('The forecasting page is separate: it uses preset or mathematically generated illustrative values, with upper and lower curves offset by 12 points and clipped to 0–100. Those curves are not statistically estimated confidence intervals, and the page does not demonstrate a trained predictive model or verified prediction accuracy. [P4]')
sub('What is proposed for NCD exposure')
p('The NCD specification proposes separate pollution, heat and cold exposure indicators. It explicitly requires scoring_enabled to remain false and scores to remain null. The dashboard currently shows historical measurements; it does not calculate disease probability or a combined NCD score. [P5]')
p('The draft pollution formula is 100 × F(x), where F(x) is the fraction of a fixed reference sample at or below daily PM2.5 concentration x. The heat formula is 100 × clip((p − 0.90) / 0.10, 0, 1); the cold formula is 100 × clip((0.10 − p) / 0.10, 0, 1). Here p is the appropriate local seasonal temperature percentile as a fraction, and clip limits the result to 0–1. For example, p = 0.97 gives 70 heat points. [P5]')
p('These are project draft proposals, not implemented or validated clinical rules. Reference years, seasonal windows, coverage criteria and calibration remain to be approved. The monthly shared CSV does not supply the required daily pollution and heat/cold inputs, and no Low/Moderate/High NCD bands are defined. [P5]')
sub('What is needed before operational use')
p('The next methodological steps are to verify data provenance and quality, make missingness explicit, calibrate thresholds and weights against suitable observed malaria outcomes, and validate on held-out times and locations. Performance, uncertainty and intended decisions should be documented and reviewed with epidemiology and public-health specialists before alerts or resource allocation rely on the score.')
page()
h('Sources and implementation references')
p('External sources below support the environmental rationale or identify data products. They do not establish EOHA’s exact scoring weights. Sources checked on 27 October 2026.')
refs=[
('[1] World Health Organization. El Niño Southern Oscillation.','Explains how temperature, rainfall and humidity affect vector suitability.','https://www.who.int/news-room/fact-sheets/detail/el-nino-southern-oscillation-%28enso%29'),
('[2] Mordecai EA and colleagues. 2013. Optimal temperature for malaria transmission is dramatically lower than previously predicted. Ecology Letters 16, 22–30. DOI 10.1111/ele.12015.','Research supporting nonlinear temperature effects; not the EOHA point schedule.','https://pubmed.ncbi.nlm.nih.gov/23050931/'),
('[3] NASA LP DAAC. MOD11A1 version 6.1. Earth Engine Data Catalog.','Defines daytime surface temperature, scaling and quality bands.','https://developers.google.com/earth-engine/datasets/catalog/MODIS_061_MOD11A1'),
('[4] NASA LP DAAC. MOD09GA version 6.1. Earth Engine Data Catalog.','Defines reflectance bands used by the exporter for NDWI and available quality flags.','https://developers.google.com/earth-engine/datasets/catalog/MODIS_061_MOD09GA'),
('[5] WorldPop. Global Project Population Data. Earth Engine Data Catalog.','Identifies estimated population counts per grid cell and the legacy collection.','https://developers.google.com/earth-engine/datasets/catalog/WorldPop_GP_100m_pop'),
('[6] ECMWF and Copernicus. ERA5-Land Hourly. Earth Engine Data Catalog.','Defines the reanalysis product used for soil moisture and air temperature.','https://developers.google.com/earth-engine/datasets/catalog/ECMWF_ERA5_LAND_HOURLY')]
for a,b,u in refs:
    p(a+' '+b+'\n'+u)
sub('Project sources')
p('[P1] scripts/malaria/dashboard/risk.js — calculate_risk_score, get_risk_label and calculate_average_risk_score; loaded by pages/malaria/dashboard.html.\n[P2] scripts/malaria/reports.js — calculateRisk, riskLabel and parsed CSV handling.\n[P3] tools/data-pipeline/export_dataset_csv.py — product selection, monthly averaging and ward aggregation.\n[P4] scripts/malaria/forecasting.js — FORECAST_DATA and illustrative bounds.\n[P5] docs/ncd/risk-scoring.md; schemas/ncd/risk-record.schema.json; scripts/ncd/data.js — proposed exposure contract and current measurement handling.')
p('Implementation baseline: local EOHA working tree reviewed on 27 October 2026, including uncommitted changes. Project paths above are relative to the repository root.')
# Reduce bibliography type slightly while retaining readable body copy.
for para in doc.paragraphs:
    if para.text.startswith(('[1]','[2]','[3]','[4]','[5]','[6]','[P1]')):
        for run in para.runs: run.font.size=Pt(10)
doc.core_properties.title='EOHA risk scoring explained'
doc.core_properties.subject='Current malaria scoring method, rationale, provenance and NCD proposal'
doc.core_properties.author='EOHA project documentation'
doc.save(out/'EOHA-risk-scoring-explained.docx')
(out/'EOHA-risk-scoring-explained.md').write_text('\n'.join(md),encoding='utf-8')
print(out/'EOHA-risk-scoring-explained.docx')
